"""
Minute exports: the PDF render contexts and the Excel workbook.

Two documents, deliberately distinct:
  * the minute PDF - the record, laid out as the manual's minute page is (p.8): date
    and time, minute type, members present and absent, the agenda table, then the
    signature blocks stamped ACKNOWLEDGED or ABSENT;
  * the acknowledgement signature sheet - one page, printed and signed in the room.

Both are allowed on an archived minute: view and export is precisely the set that
stays open once a minute becomes a record.
"""
from io import BytesIO

from . import workflow
from .models import MinuteParticipant
from .services import sanitize_minute_html


def _ordered_participants(minute):
    """
    Present first, then absent, then invitees - alphabetically within each group.

    That is the order the manual's minute page lists them in ("Members Present: …,
    Members Absent: …"), and the order the signature strip runs in.
    """
    rank = {MinuteParticipant.Attendance.PRESENT: 0,
            MinuteParticipant.Attendance.ABSENT: 1,
            MinuteParticipant.Attendance.INVITEE: 2}
    return sorted(
        minute.participants.all(),
        key=lambda p: (rank.get(p.attendance, 3), p.display_name.lower()),
    )


def build_pdf_context(minute):
    """
    Everything pdf/minute.html renders.

    The agenda body is re-sanitized here even though it was sanitized on write. This
    is the moment untrusted HTML is rendered by WeasyPrint, which can be made to fetch
    remote resources by a stylesheet the write-time gate let through; running the gate
    again at render time is cheap and closes that window.
    """
    participants = _ordered_participants(minute)
    present = [p for p in participants
               if p.attendance == MinuteParticipant.Attendance.PRESENT]
    absent = [p for p in participants
              if p.attendance == MinuteParticipant.Attendance.ABSENT]
    invitees = [p for p in participants
                if p.attendance == MinuteParticipant.Attendance.INVITEE]

    return {
        "minute": minute,
        "document_number": minute.minute_number,
        "subject": minute.display_title,
        "meeting": {
            "date": minute.meeting_date,
            "time": minute.meeting_time,
            "type": minute.minute_type.name if minute.minute_type_id else "—",
            "reference": minute.reference_number or "",
            "fro": minute.fro_name or "",
        },
        "members_present": present,
        "members_absent": absent,
        "invitee_members": invitees,
        "present_names": ", ".join(
            f"{p.display_name}({p.department_label})" if p.department_label
            else p.display_name for p in present) or "—",
        "absent_names": ", ".join(
            f"{p.display_name}({p.department_label})" if p.department_label
            else p.display_name for p in absent),
        "invitee_names": ", ".join(p.display_name for p in invitees),
        "agenda_body": sanitize_minute_html(minute.agenda_body or ""),
        "signature_rows": workflow.signature_blocks(minute),
        "ack_summary": workflow.acknowledgement_summary(minute),
        "status_label": minute.get_status_display(),
        "archived_at": minute.archived_at,
    }


def build_acknowledgement_sheet_context(minute):
    """
    The standalone acknowledgement signature sheet.

    Every member, their department, their acknowledgement status and date, and a ruled
    column to sign in. Deliberately NOT a page of the minute PDF: this is what gets
    printed and passed round a boardroom table, so it carries the minute's identity and
    nothing else.
    """
    return {
        "minute": minute,
        "document_number": minute.minute_number,
        "subject": minute.display_title,
        "meeting_date": minute.meeting_date,
        "meeting_time": minute.meeting_time,
        "meeting_type": minute.minute_type.name if minute.minute_type_id else "—",
        "participants": _ordered_participants(minute),
        "ack_summary": workflow.acknowledgement_summary(minute),
        "archived_at": minute.archived_at,
    }


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def minute_workbook(minute):
    """
    One workbook, three sheets: Minute, Members, Audit Trail.

    Three rather than the eight the previous revision wrote, because the registers the
    other five described (decisions, action items, resolutions, traceability,
    escalation) are not part of this workflow.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F3864")
    label_font = Font(bold=True)

    book = Workbook()

    def write_header(sheet, columns):
        sheet.append(columns)
        for index, _ in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=index)
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(vertical="center")

    def autosize(sheet, widths):
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width

    # --- Sheet 1: the minute itself, as label/value rows ---
    summary = book.active
    summary.title = "Minute"
    rows = [
        ("Minute Number", minute.minute_number),
        ("Reference No.", minute.reference_number or "—"),
        ("Subject", minute.display_title),
        ("Minute Type", minute.minute_type.name if minute.minute_type_id else "—"),
        ("Meeting Date", str(minute.meeting_date)),
        ("Meeting Time", minute.meeting_time.strftime("%H:%M")
         if minute.meeting_time else "—"),
        ("FRO", minute.fro_name or "—"),
        ("Initiated By", minute.created_by.get_full_name()
         or minute.created_by.username),
        ("Department", minute.department_name or "—"),
        ("Status", minute.get_status_display()),
        ("Archived At", _naive(minute.archived_at) or "—"),
    ]
    for label, value in rows:
        summary.append([label, value])
        summary.cell(row=summary.max_row, column=1).font = label_font
    autosize(summary, [22, 60])

    # --- Sheet 2: members and their acknowledgement ---
    members = book.create_sheet("Members")
    write_header(members, ["Name", "Designation", "Department", "Attendance",
                           "Acknowledgement", "Acknowledged At", "Remarks"])
    for row in _ordered_participants(minute):
        members.append([
            row.display_name,
            row.designation or "—",
            row.department_label or "—",
            row.get_attendance_display(),
            row.get_ack_status_display(),
            _naive(row.acknowledged_at) or "—",
            row.remarks or "",
        ])
    autosize(members, [26, 22, 20, 14, 22, 20, 40])

    # --- Sheet 3: the audit trail ---
    trail = book.create_sheet("Audit Trail")
    write_header(trail, ["When", "Action", "Actor", "Remarks"])
    for entry in minute.audit_entries.all():
        trail.append([
            _naive(entry.created_at),
            entry.get_action_display(),
            entry.actor_display,
            entry.remarks or "",
        ])
    autosize(trail, [20, 28, 26, 60])

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _naive(value):
    """
    openpyxl cannot write a timezone-aware datetime; it raises rather than dropping
    the offset silently. Localise then strip, so the sheet shows the wall-clock time a
    reader in Nepal expects.
    """
    if value is None:
        return None
    from django.utils import timezone
    return timezone.localtime(value).replace(tzinfo=None)
