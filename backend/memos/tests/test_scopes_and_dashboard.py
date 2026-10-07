"""
The Phase 3 sidebar menus (?scope=), the dashboard counts behind the badges, and
the Phase 6 archive export.

Every menu is defined once, on the server, as a scope. The client passes the
scope name and renders what comes back - so a menu cannot drift from its
definition, and no client re-derives who may see what.
"""
import pytest

from memos.models import Memo, MemoWorkflowStep

from .conftest import add_sections
from memos.services import generate_memo_number
from users.models import User

RoleType = MemoWorkflowStep.RoleType


def _user(username, role=User.Roles.MAKER, department="Finance"):
    return User.objects.create_user(
        username=username, email=f"{username}@nif.test", password="pass12345",
        first_name=username.capitalize(), last_name="T", role=role, department=department,
    )


def _memo(author, subject="Memo", status=Memo.Status.DRAFT, memo_type=Memo.MemoType.GENERAL):
    memo = Memo.objects.create(
        subject=subject, memo_type=memo_type, status=status, created_by=author,
        memo_number=generate_memo_number(memo_type),
    )
    add_sections(memo, ("Background", "<p>b</p>"))
    return memo


@pytest.fixture
def world(db):
    """
    One department (Finance) with an author, a department head and an approver,
    plus an outsider in another department.
    """
    people = {
        "author": _user("sc_author"),
        "colleague": _user("sc_colleague"),
        "head": _user("sc_head", User.Roles.CHECKER),
        "approver": _user("sc_approver", User.Roles.APPROVER),
        "outsider": _user("sc_outsider", User.Roles.CHECKER, department="Engineering"),
    }
    return people


def _scope_ids(api, user, scope):
    api.force_authenticate(user)
    res = api.get("/api/v1/memos/", {"scope": scope})
    assert res.status_code == 200, res.data
    return {row["id"] for row in res.data["results"]}


def _send(api, author, memo, rows):
    api.force_authenticate(author)
    res = api.post(f"/api/v1/memos/{memo.id}/send-for-review/", {"workflow": rows}, format="json")
    assert res.status_code == 200, res.data
    return res


@pytest.mark.django_db
def test_drafts_scope_returns_only_my_own_drafts(api, world):
    mine = _memo(world["author"], "My draft")
    _memo(world["colleague"], "Someone else's draft")
    assert _scope_ids(api, world["author"], "drafts") == {str(mine.id)}


@pytest.mark.django_db
def test_draft_for_review_and_outbox_track_what_i_sent(api, world):
    memo = _memo(world["author"], "Sent for review")
    _send(api, world["author"], memo, [
        {"assignee_id": str(world["head"].id), "role_type": RoleType.REVIEWER},
        {"assignee_id": str(world["approver"].id), "role_type": RoleType.APPROVER},
    ])
    still_draft = _memo(world["author"], "Still a draft")

    assert _scope_ids(api, world["author"], "draft_for_review") == {str(memo.id)}
    assert _scope_ids(api, world["author"], "outbox") == {str(memo.id)}
    assert _scope_ids(api, world["author"], "drafts") == {str(still_draft.id)}


@pytest.mark.django_db
def test_pending_scope_is_only_what_is_waiting_on_me_now(api, world):
    memo = _memo(world["author"], "Two-stage memo")
    _send(api, world["author"], memo, [
        {"assignee_id": str(world["head"].id), "role_type": RoleType.REVIEWER},
        {"assignee_id": str(world["approver"].id), "role_type": RoleType.APPROVER},
    ])

    # The reviewer's turn is now; the approver's is not.
    assert _scope_ids(api, world["head"], "pending") == {str(memo.id)}
    assert _scope_ids(api, world["approver"], "pending") == set()

    # But the approver can already see it in their inbox, so they can read the
    # memo and the reviewer's remarks before their turn arrives.
    assert str(memo.id) in _scope_ids(api, world["approver"], "inbox")

    api.force_authenticate(world["head"])
    api.post(f"/api/v1/memos/{memo.id}/act/",
             {"decision": "proceed", "remarks": "Checked, all in order."}, format="json")

    assert _scope_ids(api, world["head"], "pending") == set()
    assert _scope_ids(api, world["approver"], "pending") == {str(memo.id)}


@pytest.mark.django_db
def test_inbox_excludes_my_own_memos(api, world):
    """Inbox is incoming work; my own memo belongs in the outbox."""
    memo = _memo(world["author"], "Mine")
    _send(api, world["author"], memo, [
        {"assignee_id": str(world["approver"].id), "role_type": RoleType.APPROVER},
    ])
    assert _scope_ids(api, world["author"], "inbox") == set()
    assert _scope_ids(api, world["author"], "outbox") == {str(memo.id)}


@pytest.mark.django_db
def test_department_scope_is_bounded_by_department(api, world):
    finance = _memo(world["author"], "Finance memo", status=Memo.Status.DRAFT_FOR_REVIEW)
    engineering = _memo(
        _user("eng_author", department="Engineering"), "Engineering memo",
        status=Memo.Status.DRAFT_FOR_REVIEW,
    )

    head_view = _scope_ids(api, world["head"], "department")
    assert str(finance.id) in head_view
    assert str(engineering.id) not in head_view

    # A department head elsewhere sees their own unit's memo, not Finance's.
    outsider_view = _scope_ids(api, world["outsider"], "department")
    assert str(engineering.id) in outsider_view
    assert str(finance.id) not in outsider_view


@pytest.mark.django_db
def test_approved_and_archived_scopes(api, world):
    memo = _memo(world["author"], "To be approved")
    _send(api, world["author"], memo, [
        {"assignee_id": str(world["approver"].id), "role_type": RoleType.APPROVER},
    ])
    api.force_authenticate(world["approver"])
    api.post(f"/api/v1/memos/{memo.id}/act/", {"decision": "proceed", "remarks": "Checked the memo and its attachments; content is in order."}, format="json")

    memo.refresh_from_db()
    assert memo.status == Memo.Status.ARCHIVED

    # Approval auto-archives, so "Approved" spans both states or it would read
    # empty for every completed memo.
    assert _scope_ids(api, world["author"], "approved") == {str(memo.id)}
    assert _scope_ids(api, world["author"], "archived") == {str(memo.id)}


@pytest.mark.django_db
def test_dashboard_counts_the_whole_visible_set_not_one_page(api, world):
    """
    The old dashboard card counted rows inside a paginated list response, so it
    under-reported as soon as a user could see more than one page (PAGE_SIZE=50).
    """
    for index in range(55):
        _memo(world["author"], f"Draft {index}")

    api.force_authenticate(world["author"])
    res = api.get("/api/v1/memos/dashboard/")
    assert res.status_code == 200
    assert res.data["drafts"] == 55
    assert res.data["my_memos"] == 55

    listed = api.get("/api/v1/memos/", {"scope": "drafts"})
    assert len(listed.data["results"]) == 50  # one page
    assert listed.data["count"] == 55         # but the count is honest


@pytest.mark.django_db
def test_dashboard_pending_and_inbox_counts(api, world):
    memo = _memo(world["author"], "Needs two approvals")
    _send(api, world["author"], memo, [
        {"assignee_id": str(world["head"].id), "role_type": RoleType.REVIEWER},
        {"assignee_id": str(world["approver"].id), "role_type": RoleType.APPROVER},
    ])

    api.force_authenticate(world["head"])
    head = api.get("/api/v1/memos/dashboard/").data
    assert head["pending_actions"] == 1
    assert head["inbox"] == 1

    api.force_authenticate(world["approver"])
    approver = api.get("/api/v1/memos/dashboard/").data
    assert approver["pending_actions"] == 0
    assert approver["inbox"] == 1


@pytest.mark.django_db
def test_export_returns_an_xlsx_of_the_current_scope(api, world):
    memo = _memo(world["author"], "Exportable memo")
    _send(api, world["author"], memo, [
        {"assignee_id": str(world["approver"].id), "role_type": RoleType.APPROVER},
    ])
    api.force_authenticate(world["approver"])
    api.post(f"/api/v1/memos/{memo.id}/act/", {"decision": "proceed", "remarks": "Checked the memo and its attachments; content is in order."}, format="json")

    api.force_authenticate(world["author"])
    res = api.get("/api/v1/memos/export/", {"scope": "archived"})
    assert res.status_code == 200
    assert "spreadsheetml" in res["Content-Type"]
    assert "archived-export.xlsx" in res["Content-Disposition"]

    import io
    from openpyxl import load_workbook
    sheet = load_workbook(io.BytesIO(res.content)).active
    rows = list(sheet.values)
    assert rows[0][0] == "Memo Number"
    # Look columns up by header rather than by index, so adding a column to the
    # export does not silently shift every assertion onto the wrong field.
    column = {label: index for index, label in enumerate(rows[0])}
    row = rows[1]
    assert row[column["Memo Number"]] == memo.memo_number
    assert row[column["Subject"]] == "Exportable memo"
    assert row[column["Status"]] == "Archived"
    assert row[column["Department"]] == "Finance"
    assert row[column["Approver"]] == "Sc_approver T"


@pytest.mark.django_db
def test_export_cannot_leak_memos_the_caller_cannot_read(api, world):
    """Export runs the list queryset, so scoping is inherited, not re-derived."""
    import io
    from openpyxl import load_workbook

    _memo(world["colleague"], "Private to a colleague")
    mine = _memo(world["author"], "Mine")

    api.force_authenticate(world["author"])
    res = api.get("/api/v1/memos/export/")
    sheet = load_workbook(io.BytesIO(res.content)).active
    numbers = [row[0] for row in list(sheet.values)[1:]]
    assert numbers == [mine.memo_number]


# ---------------------------------------------------------------------------
# Phase 4 employee directory
# ---------------------------------------------------------------------------
@pytest.mark.django_db
def test_employee_directory_returns_any_employee_regardless_of_role(api, world):
    """
    Phase 4 is explicit that any employee may be selected for any role type, so
    unlike available-checkers this directory is not role-filtered.
    """
    api.force_authenticate(world["author"])
    res = api.get("/api/v1/memos/employees/", {"search": "sc_"})
    assert res.status_code == 200
    roles = {row["role"] for row in res.data}
    assert {"maker", "checker", "approver"} <= roles
    # Never the caller themselves, and never an email address.
    assert str(world["author"].id) not in {row["id"] for row in res.data}
    assert all("email" not in row for row in res.data)


@pytest.mark.django_db
def test_employee_directory_requires_a_search_term(api, world):
    """Search-gated so it cannot be walked to dump the staff roster (H5)."""
    api.force_authenticate(world["author"])
    assert api.get("/api/v1/memos/employees/").data == []
    assert api.get("/api/v1/memos/employees/", {"search": "s"}).data == []
    assert api.get("/api/v1/memos/employees/", {"search": "sc_head"}).data != []


@pytest.mark.django_db
def test_employee_directory_matches_designation_and_employee_id(api, world):
    person = world["head"]
    person.designation = "Chief Engineer"
    person.save(update_fields=["designation"])

    api.force_authenticate(world["author"])
    res = api.get("/api/v1/memos/employees/", {"search": "Chief Eng"})
    assert [row["id"] for row in res.data] == [str(person.id)]
    assert res.data[0]["designation"] == "Chief Engineer"
    assert res.data[0]["department"] == "Finance"


# ---------------------------------------------------------------------------
# The manual's dashboard (E-memo-manual p.2)
# ---------------------------------------------------------------------------
@pytest.mark.django_db
class TestTheManualsDashboardTiles:
    """
    Four tiles, and the manual defines each one. The two "Under Process" figures
    answer different questions and must not be the same number.
    """

    def test_initiated_counts_what_i_raised_and_is_still_moving(self, api, world):
        mine = _memo(world["author"], "Mine, in flight")
        _send(api, world["author"], mine, [
            {"assignee_id": str(world["approver"].id), "role_type": RoleType.APPROVER},
        ])
        _memo(world["author"], "Mine, still a draft")

        api.force_authenticate(world["author"])
        counts = api.get("/api/v1/memos/dashboard/").data
        assert counts["under_process_initiated"] == 1, "a draft is not under process"
        assert counts["drafts"] == 1

    def test_involvement_counts_what_i_have_acted_on_not_what_awaits_me(
            self, api, world):
        """
        "Memo in which the logged-in user has been involved (i.e. Supported,
        Reviewed, Approved, and Noted)" — so a memo merely sitting on my desk does
        not count, and one I have already acted on does.
        """
        memo = _memo(world["author"], "Two-step")
        _send(api, world["author"], memo, [
            {"assignee_id": str(world["head"].id), "role_type": RoleType.REVIEWER},
            {"assignee_id": str(world["approver"].id), "role_type": RoleType.APPROVER},
        ])

        # It is waiting on the reviewer, who has done nothing yet.
        api.force_authenticate(world["head"])
        assert api.get("/api/v1/memos/dashboard/").data[
            "under_process_involvement"] == 0

        # Once they act it is theirs by involvement - and it is still in flight,
        # because the approver has not acted.
        # A reviewer's action requires a comment; the workflow enforces it.
        acted = api.post(f"/api/v1/memos/{memo.id}/act/",
                         {"decision": "proceed",
                          "remarks": "Reviewed and in order."}, format="json")
        assert acted.status_code == 200, acted.data
        assert api.get("/api/v1/memos/dashboard/").data[
            "under_process_involvement"] == 1

    def test_involvement_excludes_my_own_memos(self, api, world):
        """An author's own memo belongs to Initiated; counting it twice would make
        the two tiles overlap."""
        memo = _memo(world["author"], "Mine")
        _send(api, world["author"], memo, [
            {"assignee_id": str(world["approver"].id), "role_type": RoleType.APPROVER},
        ])
        api.force_authenticate(world["author"])
        counts = api.get("/api/v1/memos/dashboard/").data
        assert counts["under_process_initiated"] == 1
        assert counts["under_process_involvement"] == 0
