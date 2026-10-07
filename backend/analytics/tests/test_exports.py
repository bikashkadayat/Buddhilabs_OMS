"""The five analytics exports.

Two things are proved here. First that each builder produces a file that
actually opens -- a PDF nobody can open is a failure that a 200 response hides.
Second, and more important, that an exported figure MATCHES the dashboard
endpoint for the same window. They call the same metric function by design, and
this is the test that keeps it that way.
"""
import csv
import io

import pytest
from django.urls import reverse

from analytics import exports

pytestmark = pytest.mark.django_db


@pytest.fixture
def busy_month(org, last_month_days, mark, take_leave, comp_entry):
    from attendance.models import Attendance

    for member in org["eng"]:
        mark(member, last_month_days[:-1], overtime_hours="1.50")
    for member in org["ops"]:
        mark(member, last_month_days[:3], status=Attendance.Status.LATE,
             late_minutes=20)
    take_leave(org["eng"][0], last_month_days[-1:])
    comp_entry(org["eng"][0], "1.00", source_date=last_month_days[0])
    return org


def params_for(last_month, fmt):
    start, end = last_month
    return {"from": start.isoformat(), "to": end.isoformat(), "format": fmt}


class TestBuilders:
    def test_executive_summary_is_a_pdf(self, busy_month, last_month):
        content, filename, content_type = exports.build_executive_summary(
            params_for(last_month, "pdf"))
        assert content[:4] == b"%PDF"
        assert filename.endswith(".pdf")
        assert content_type == "application/pdf"

    def test_department_analytics_is_a_pdf(self, busy_month, last_month):
        content, filename, _ct = exports.build_department_analytics(
            params_for(last_month, "pdf"))
        assert content[:4] == b"%PDF"
        assert "department_analytics" in filename

    def test_attendance_analytics_workbook_opens_with_four_sheets(
            self, busy_month, last_month):
        from openpyxl import load_workbook

        content, filename, content_type = exports.build_attendance_analytics(
            params_for(last_month, "excel"))
        workbook = load_workbook(io.BytesIO(content))
        assert workbook.sheetnames == ["Summary", "Trend", "Departments", "Comparisons"]
        assert filename.endswith(".xlsx")
        assert "spreadsheetml" in content_type

    def test_overtime_analytics_workbook_opens(self, busy_month, last_month):
        from openpyxl import load_workbook

        content, _filename, _ct = exports.build_overtime_analytics(
            params_for(last_month, "excel"))
        workbook = load_workbook(io.BytesIO(content))
        assert workbook.sheetnames == ["Summary", "By department", "Trend"]

    def test_workforce_analytics_csv_parses(self, busy_month, last_month):
        content, filename, content_type = exports.build_workforce_analytics(
            params_for(last_month, "csv"))
        rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
        assert rows[0][0] == "Window start"
        assert len(rows) > 1
        assert filename.endswith(".csv")
        assert content_type == "text/csv"


class TestNumbersMatchTheDashboard:
    def test_workforce_csv_matches_the_departments_endpoint(
            self, busy_month, last_month, auth):
        """One formula, two consumers. If this drifts, an export is lying."""
        start, end = last_month
        response = auth(busy_month["admin"]).get(reverse("analytics-departments"), {
            "from": start.isoformat(), "to": end.isoformat()})
        api_rows = {row["department"]: row
                    for row in response.data["data"]["departments"]}

        content, _filename, _ct = exports.build_workforce_analytics(
            params_for(last_month, "csv"))
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        exported = {row["Department"]: row for row in reader}

        assert set(exported) == set(api_rows)
        for name, row in exported.items():
            expected = api_rows[name]["compliance_pct"]
            actual = float(row["Compliance %"]) if row["Compliance %"] else None
            assert actual == expected, f"{name} compliance drifted"

    def test_attendance_workbook_matches_the_attendance_endpoint(
            self, busy_month, last_month, auth):
        from openpyxl import load_workbook

        start, end = last_month
        response = auth(busy_month["admin"]).get(reverse("analytics-attendance"), {
            "from": start.isoformat(), "to": end.isoformat(),
            "granularity": "month"})
        kpis = response.data["data"]["kpis"]

        content, _filename, _ct = exports.build_attendance_analytics(
            params_for(last_month, "excel"))
        sheet = load_workbook(io.BytesIO(content))["Summary"]
        values = {row[0]: row[1] for row in sheet.iter_rows(min_row=2, values_only=True)}

        assert values["Expected working days"] == kpis["expected_days"]
        assert values["Days attended"] == kpis["attended_days"]
        assert values["Overtime hours"] == kpis["overtime_hours"]


class TestScopedExports:
    def test_employee_ids_narrow_the_export(self, busy_month, last_month):
        """This is how a department head's export stays department-scoped."""
        engineering = [str(member.pk) for member in busy_month["eng"]]
        params = params_for(last_month, "csv")
        params["employee_ids"] = engineering

        content, _filename, _ct = exports.build_workforce_analytics(params)
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        departments = {row["Department"] for row in reader}
        assert departments == {"Engineering"}

    def test_export_never_names_an_individual(self, busy_month, last_month):
        """The phase's hardest constraint, enforced on the file itself."""
        content, _filename, _ct = exports.build_workforce_analytics(
            params_for(last_month, "csv"))
        text = content.decode("utf-8-sig")
        for member in busy_month["eng"] + busy_month["ops"]:
            assert member.get_full_name() not in text
            assert str(member.pk) not in text


class TestThroughTheReportsHub:
    def test_generation_end_to_end_produces_a_ready_run(self, busy_month, auth,
                                                        settings, last_month):
        """Proves the builders plug into the Phase 8 machinery unchanged."""
        settings.REPORTS_RUN_SYNC = True
        from reports.models import ReportRun

        start, end = last_month
        response = auth(busy_month["admin"]).post(
            reverse("report-request-report"),
            {"report_type": "executive_summary",
             "params": {"format": "pdf", "from": start.isoformat(),
                        "to": end.isoformat()}},
            format="json")
        assert response.status_code == 202

        run = ReportRun.objects.get(pk=response.data["id"])
        assert run.status == ReportRun.Status.READY, run.error
        assert run.file

    @pytest.mark.parametrize("report_type,fmt", [
        ("department_analytics", "pdf"),
        ("attendance_analytics", "excel"),
        ("overtime_analytics", "excel"),
        ("workforce_analytics", "csv"),
    ])
    def test_every_analytics_type_generates(self, report_type, fmt, busy_month,
                                            auth, settings, last_month):
        settings.REPORTS_RUN_SYNC = True
        from reports.models import ReportRun

        start, end = last_month
        response = auth(busy_month["admin"]).post(
            reverse("report-request-report"),
            {"report_type": report_type,
             "params": {"format": fmt, "from": start.isoformat(),
                        "to": end.isoformat()}},
            format="json")
        run = ReportRun.objects.get(pk=response.data["id"])
        assert run.status == ReportRun.Status.READY, run.error


class TestEmptyData:
    def test_builders_survive_an_organisation_with_no_attendance(self, org, last_month):
        """An empty window must produce an empty document, not a traceback."""
        for builder, fmt in ((exports.build_executive_summary, "pdf"),
                             (exports.build_department_analytics, "pdf"),
                             (exports.build_attendance_analytics, "excel"),
                             (exports.build_overtime_analytics, "excel"),
                             (exports.build_workforce_analytics, "csv")):
            content, _filename, _ct = builder(params_for(last_month, fmt))
            assert content
